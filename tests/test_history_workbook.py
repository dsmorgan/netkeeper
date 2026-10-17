"""Reading the old mailing tool's workbook (#65, Part A): hand-built fixtures only."""

from __future__ import annotations

import re
import time
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from io import BytesIO

import pytest
from history_fixtures import Tab, grid, workbook_bytes
from openpyxl import Workbook, load_workbook

from netkeeper.crm.history_workbook import (
    BLANK_ROWS_END_A_LIST,
    MAX_SHEET_COLUMNS,
    Person,
    WorkbookError,
    read_tab,
    read_workbook,
)


def test_a_tab_reads_its_header_and_three_lists() -> None:
    campaign = read_tab(0, grid(Tab()))
    assert campaign.name == "Spring check-in"
    assert campaign.subject == "Catching up"
    assert campaign.started_on == date(2026, 3, 2)
    assert campaign.last_batch_on == date(2026, 3, 9)
    assert (campaign.recipients_count, campaign.opens_count, campaign.bounces_count) == (5, 3, 1)
    assert [p.email for p in campaign.opened] == [
        "ada@example.test",
        "bob@example.test",  # lowercased
        "cy@example.test",
    ]
    # The clicked link under each person is not a person.
    assert campaign.clicked == (Person("Ada", "Lovelace", "ada@example.test"),)
    assert campaign.bounced == (Person("Dee", "Turing", "dee@example.test"),)
    assert campaign.warnings == ()


def test_recipients_are_everyone_listed_once_with_the_lists_that_named_them() -> None:
    campaign = read_tab(0, grid(Tab()))
    by_email = {r.email: r for r in campaign.recipients}
    assert list(by_email) == [
        "ada@example.test",
        "bob@example.test",
        "cy@example.test",
        "dee@example.test",
    ]
    ada, dee = by_email["ada@example.test"], by_email["dee@example.test"]
    assert (ada.opened, ada.clicked, ada.bounced) == (True, True, False)
    assert (dee.opened, dee.clicked, dee.bounced) == (False, False, True)
    # The Recipients cell counts five; the lists name four.
    assert campaign.unlisted == 1


@pytest.mark.parametrize("clicks_col", [11, 10])
def test_both_column_offsets_of_the_clicks_list_read_the_same(clicks_col: int) -> None:
    clicked = (
        ("Ada", "Lovelace", "ada@example.test"),
        ("Eve", "Noether", "eve@example.test"),
    )
    campaign = read_tab(0, grid(Tab(clicks_col=clicks_col, clicked=clicked)))
    assert [p.email for p in campaign.clicked] == ["ada@example.test", "eve@example.test"]
    assert [p.email for p in campaign.opened] == [
        "ada@example.test",
        "bob@example.test",
        "cy@example.test",
    ]


def test_excel_serial_dates_are_converted() -> None:
    # 46083 is 2026-03-02 in Excel's 1900 date system; .5 is noon.
    campaign = read_tab(0, grid(Tab(started=46083.5, last_batch=46090)))
    assert campaign.started_on == date(2026, 3, 2)
    assert campaign.last_batch_on == date(2026, 3, 9)


def test_iso_text_dates_are_read() -> None:
    campaign = read_tab(0, grid(Tab(started="2026-03-02 09:30", last_batch="2026-03-09")))
    assert (campaign.started_on, campaign.last_batch_on) == (date(2026, 3, 2), date(2026, 3, 9))


def test_a_short_list_and_an_empty_clicks_list() -> None:
    tab = Tab(
        opened=(("Ada", "Lovelace", "ada@example.test"),),
        opens="1 (100.0%)",
        recipients=1,
        clicked=(),
        bounces=(),
        bounced=0,
    )
    campaign = read_tab(0, grid(tab))
    assert [p.email for p in campaign.opened] == ["ada@example.test"]
    assert campaign.clicked == ()
    assert campaign.bounced == ()
    assert campaign.unlisted == 0
    assert campaign.warnings == ()


def test_a_list_ends_at_two_blank_rows_in_a_row() -> None:
    assert BLANK_ROWS_END_A_LIST == 2
    # One blank row inside the clicks list is skipped; two end it.
    far = 3 + 3 * 1 + 2  # past the one click group and two blank rows
    tab = Tab(extra_rows=[(far, 11, "Zed"), (far, 12, "Far"), (far, 13, "zed@example.test")])
    campaign = read_tab(0, grid(tab))
    assert [p.email for p in campaign.clicked] == ["ada@example.test"]


def test_a_row_with_a_name_but_no_address_is_counted_as_a_warning() -> None:
    tab = Tab(
        opened=(("Ada", "Lovelace", "ada@example.test"), ("Bob", "Babbage", "not an address")),
        opens="2 (40.0%)",
    )
    campaign = read_tab(0, grid(tab))
    assert [p.email for p in campaign.opened] == ["ada@example.test"]
    assert campaign.warnings == (
        "1 row(s) of the opened list were dropped: no usable address",
        "the opens list has 1 people but the opens cell says 2",
    )


def test_only_a_link_shaped_row_in_the_clicks_list_is_a_link() -> None:
    """A text-only row is a clicked link only in the clicks list, and only when it looks
    like one; anything else is reported (#65 review, N3)."""
    tab = Tab(
        extra_rows=[
            (5, 13, "a note, not a link"),  # the blank row after Ada's link group
            (8, 8, "stray text in the opens list"),  # right below the three openers
        ]
    )
    campaign = read_tab(0, grid(tab))
    assert [p.email for p in campaign.clicked] == ["ada@example.test"]
    assert sorted(campaign.warnings) == [
        "1 row(s) of the clicked list were dropped: no usable address",
        "1 row(s) of the opened list were dropped: no usable address",
    ]


@pytest.mark.parametrize("value", ["mailto:ada@example.test", "ada:x@example.test"])
def test_a_colon_in_the_local_part_is_not_an_address(value: str) -> None:
    tab = Tab(
        opened=(("Ada", "Lovelace", value), ("Bob", "Babbage", "bob@example.test")),
        opens="2 (40.0%)",
    )
    campaign = read_tab(0, grid(tab))
    assert [p.email for p in campaign.opened] == ["bob@example.test"]


def test_a_formula_with_no_saved_value_is_reported_not_read_as_blank() -> None:
    tab = Tab(
        opened=(
            ("Ada", "Lovelace", '=LOWER("ADA@example.test")'),
            ("Bob", "Babbage", "bob@example.test"),
        ),
        opens="2 (40.0%)",
    )
    (campaign,) = read_workbook(workbook_bytes(tab)).campaigns
    assert campaign.warnings[0] == (
        "1 formula cell(s) have no saved value and read as empty;"
        " open the sheet, let it calculate, and export it again"
    )
    assert "1 row(s) of the opened list were dropped: no usable address" in campaign.warnings


def test_a_tab_without_a_campaign_header_is_refused() -> None:
    with pytest.raises(WorkbookError, match="Campaign name"):
        read_tab(0, [["Summary"], [None]])


def test_a_tab_without_a_start_date_is_refused() -> None:
    with pytest.raises(WorkbookError, match="start date"):
        read_tab(0, grid(Tab(started=None)))


def test_a_workbook_reads_every_tab_and_skips_one_that_is_not_a_campaign() -> None:
    data = workbook_bytes(
        Tab(title="First"),
        Tab(title="Second", name="Follow-up", started=datetime(2026, 3, 16), clicks_col=10),
    )
    workbook = read_workbook(data)
    assert [c.name for c in workbook.campaigns] == ["Spring check-in", "Follow-up"]
    assert workbook.campaigns[1].started_on == date(2026, 3, 16)
    assert len(workbook.sha256) == 64
    assert workbook.skipped == ()


def test_a_tab_that_is_not_a_campaign_is_skipped_with_a_reason() -> None:
    book = load_workbook(BytesIO(workbook_bytes(Tab())))
    book.create_sheet("Summary")["A1"] = "Totals"
    buffer = BytesIO()
    book.save(buffer)
    workbook = read_workbook(buffer.getvalue())
    assert [c.name for c in workbook.campaigns] == ["Spring check-in"]
    assert [(s.index, s.reason) for s in workbook.skipped] == [(1, "no 'Campaign name' header")]


def test_a_file_that_is_not_a_workbook_is_refused() -> None:
    with pytest.raises(WorkbookError):
        read_workbook(b"First Name,Last Name,Email\n")


@pytest.mark.parametrize(("row", "col"), [(8, 8), (9, 3)])
def test_a_text_only_row_outside_the_clicks_list_is_reported_not_a_link(row: int, col: int) -> None:
    """N25: only the clicks list holds links; link-shaped text elsewhere is a dropped row."""
    campaign = read_tab(0, grid(Tab(extra_rows=[(row, col, "example.test/landing page")])))
    assert len(campaign.warnings) == 1
    assert campaign.warnings[0].endswith("list were dropped: no usable address")


def test_a_malformed_address_in_the_clicks_list_is_reported_not_taken_for_a_link() -> None:
    """Nit 6: no name, and an address with a stray character."""
    campaign = read_tab(0, grid(Tab(extra_rows=[(6, 13, "eve@example.test;")])))
    assert campaign.warnings == ("1 row(s) of the clicked list were dropped: no usable address",)


def test_a_personalized_subject_is_warned_about() -> None:
    """Nit 4: the scan's exact subject match cannot work on {{first_name}}."""
    campaign = read_tab(0, grid(Tab(subject="Catching up, {{first_name}}?")))
    assert campaign.warnings == (
        "the subject is personalized ({{...}}), so the Gmail scan's exact subject"
        " match cannot find replies to it",
    )


# --- guards (#619) -----------------------------------------------------------------

_TIME_BOUND_SECONDS = 10.0


def _cells_workbook(cells: Sequence[tuple[int, int]], *, dimension: bool = True) -> bytes:
    """A workbook with a ``"x"`` in each (row, column); optionally without its
    ``<dimension>`` element, the way some writers leave it out."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row, col in cells:
        sheet.cell(row=row, column=col, value="x")
    buffer = BytesIO()
    book.save(buffer)
    if dimension:
        return buffer.getvalue()
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(buffer.getvalue())) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload = re.sub(rb"<dimension[^>]*/>", b"", payload)
            dst.writestr(info, payload, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def _refused(data: bytes, match: str) -> None:
    started = time.monotonic()
    with pytest.raises(WorkbookError, match=match):
        read_workbook(data)
    assert time.monotonic() - started < _TIME_BOUND_SECONDS


def test_a_far_away_cell_is_refused_on_its_declared_dimension() -> None:
    _refused(_cells_workbook([(1_048_576, 1)]), "row limit")
    _refused(_cells_workbook([(1, 16_384)]), "column limit")


def test_a_far_away_cell_is_refused_while_reading_when_there_is_no_dimension() -> None:
    data = _cells_workbook([(1_048_576, 1)], dimension=False)
    assert b"<dimension" not in zipfile.ZipFile(BytesIO(data)).read("xl/worksheets/sheet1.xml")
    _refused(data, "more than 50000 rows")
    _refused(_cells_workbook([(1, 16_384)], dimension=False), "64 limit")


def test_a_dense_padded_sheet_is_refused_on_its_cell_count() -> None:
    # One cell per row in the last allowed column: small on disk, but every row is padded
    # to 64 cells when read. 20,000 rows x 64 = 1.28M cells, over the 1M cap, though each
    # axis is inside its own limit.
    cells = [(row, MAX_SHEET_COLUMNS) for row in range(1, 20_001)]
    _refused(_cells_workbook(cells, dimension=False), "more than 1000000 cells")
    _refused(_cells_workbook(cells), "cells declared")


def test_a_zip_bomb_workbook_is_refused() -> None:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (80 * 1024 * 1024))
    _refused(buffer.getvalue(), "100x ratio")


def test_a_normal_export_still_reads() -> None:
    book = read_workbook(workbook_bytes(Tab()))
    assert len(book.campaigns) == 1
    assert book.campaigns[0].name == "Spring check-in"
