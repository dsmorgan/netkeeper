"""Hand-built workbooks shaped like the old mailing tool's report (#65).

Every name and address here is invented. The geometry follows the real report's:
a header row of labels with the values below it; an opens list under a ``First
Name`` header beside a note; a bounce list below the note, with a blank row after
its header; a clicks list under a ``Clicks`` label, each person followed by the
link they clicked and a blank row. ``clicks_col`` moves the clicks list, as the
report does on some tabs.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

Person = tuple[str, str, str]


@dataclass
class Tab:
    name: str = "Spring check-in"
    subject: str = "Catching up"
    started: object = datetime(2026, 3, 2, 9, 30)
    last_batch: object = datetime(2026, 3, 9, 9, 30)
    status: str = "Finished"
    recipients: object = 5.0
    opens: object = "3 (60.0%)"
    bounced: object = 1.0
    opened: Sequence[Person] = (
        ("Ada", "Lovelace", "ada@example.test"),
        ("Bob", "Babbage", "Bob@Example.test"),
        ("Cy", "Hopper", "cy@example.test"),
    )
    clicked: Sequence[Person] = (("Ada", "Lovelace", "ada@example.test"),)
    bounces: Sequence[Person] = (("Dee", "Turing", "dee@example.test"),)
    clicks_col: int = 11
    title: str = "Tab"
    extra_rows: list[tuple[int, int, object]] = field(default_factory=list)


def tab_rows(tab: Tab) -> dict[tuple[int, int], object]:
    """The tab's cells by one-based (row, column)."""
    cells: dict[tuple[int, int], object] = {}
    labels = (
        "Campaign Name",
        "Email Subject",
        "Start date",
        "Last Batch Sent",
        "Status",
        "Recipients",
        "Opens",
        "Bounced",
    )
    values = (
        tab.name,
        tab.subject,
        tab.started,
        tab.last_batch,
        tab.status,
        tab.recipients,
        tab.opens,
        tab.bounced,
    )
    for col, (label, value) in enumerate(zip(labels, values, strict=True), start=1):
        cells[(1, col)] = label
        cells[(2, col)] = value
    # The opens list, beside a note.
    cells[(4, 1)] = "The people below opened the email at least once."
    for col, label in enumerate(("First Name", "Last Name", "Email"), start=6):
        cells[(4, col)] = label
    for row, person in enumerate(tab.opened, start=5):
        for col, value in enumerate(person, start=6):
            cells[(row, col)] = value
    # The bounce list: a blank row after its header.
    for col, label in enumerate(("First Name", "Last Name", "Bounced Email"), start=1):
        cells[(6, col)] = label
    for row, person in enumerate(tab.bounces, start=8):
        for col, value in enumerate(person, start=1):
            cells[(row, col)] = value
    # The clicks list: person, link, blank.
    c = tab.clicks_col
    cells[(1, c)] = "Clicks"
    for offset, label in enumerate(("First Name", "Last Name", "Email")):
        cells[(2, c + offset)] = label
    row = 3
    for person in tab.clicked:
        for offset, value in enumerate(person):
            cells[(row, c + offset)] = value
        cells[(row + 1, c + 2)] = "example.test/landing page"
        row += 3
    for r, col, value in tab.extra_rows:
        cells[(r, col)] = value
    return cells


def grid(tab: Tab) -> list[list[object]]:
    """The tab as ``read_tab`` takes it: zero-based rows of values."""
    cells = tab_rows(tab)
    height = max(r for r, _ in cells)
    width = max(c for _, c in cells)
    return [[cells.get((r, c)) for c in range(1, width + 1)] for r in range(1, height + 1)]


def workbook_bytes(*tabs: Tab) -> bytes:
    book = Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    for tab in tabs:
        sheet = book.create_sheet(tab.title)
        for (row, col), value in tab_rows(tab).items():
            sheet.cell(row=row, column=col, value=value)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
