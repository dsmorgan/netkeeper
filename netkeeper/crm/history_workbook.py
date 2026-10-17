"""Read the old mailing tool's workbook: one tab per campaign (#65, Part A).

Each tab is a formatted report, not a table. It has a header block and
per-person lists side by side:

- **Header.** A row of labels (``Campaign name``, an email subject, a start
  date, ``Last batch sent``, ``Status``, ``Recipients``, ``Opens``,
  ``Bounced``) with the values in the row below. ``Recipients`` and ``Bounced``
  are numbers; ``Opens`` is text such as ``78 (78.0%)``.
- **Per-person lists.** Each starts at a ``First Name`` header cell, with
  ``Last Name`` and an address column to its right, and runs down until two
  empty rows in a row. A blank row inside a list is skipped. The list under a
  ``Clicks`` label is the clicks list, and in it a row that holds only text in
  the address column is the link the person above clicked. The list whose
  address column header names a bounce is the bounce list. The remaining one is
  the opens list: its length matches the ``Opens`` count, and it does **not**
  name everyone the campaign was sent to.

Column offsets are read from the header cells, never assumed: they differ
between tabs. Dates are ``datetime`` cells, Excel serial numbers (converted),
or ISO text.

Nothing here touches the database, and nothing here logs a value from the file:
a warning names a tab by its position and a list by its kind, never a name or an
address.
"""

from __future__ import annotations

import enum
import hashlib
import logging
import re
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any, Final

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from netkeeper.zipguard import ZipRefusal, guard_zip_members

log = logging.getLogger(__name__)

#: A list ends at this many empty rows in a row.
BLANK_ROWS_END_A_LIST: Final = 2

_ADDRESS: Final = re.compile(r"[^@\s<>,;:\"']+@[^@\s<>,;:\"']+\.[^@\s<>,;:\"']+")
_LINK: Final = re.compile(
    r"(?:https?://|\b[a-z0-9-]+(?:\.[a-z0-9-]+)*\.[a-z]{2,}\b)", re.IGNORECASE
)
_LEADING_INT: Final = re.compile(r"\s*(\d+)")
_SPACES: Final = re.compile(r"\s+")


#: A sheet over these limits is refused. A real campaign tab uses about 15 columns (a header
#: block and three lists side by side) and as many rows as its longest list. The limits apply
#: to the declared dimension and again to what is actually read, since a sheet with no
#: dimension element, or a wrong one, is bounded only by the second.
MAX_SHEET_ROWS: Final = 50_000
MAX_SHEET_COLUMNS: Final = 64
#: Counted on the cells read, after openpyxl pads every row to the sheet's width.
MAX_SHEET_CELLS: Final = 1_000_000


class WorkbookError(ValueError):
    """The file is not a workbook this reader understands."""


class ListKind(enum.StrEnum):
    OPENED = "opened"
    CLICKED = "clicked"
    BOUNCED = "bounced"


@dataclass(frozen=True, slots=True)
class Person:
    """One row of a per-person list. ``email`` is trimmed and lowercased."""

    first_name: str | None
    last_name: str | None
    email: str


@dataclass(frozen=True, slots=True)
class Recipient:
    """One person the tab names, with the lists that named them."""

    first_name: str | None
    last_name: str | None
    email: str
    opened: bool
    clicked: bool
    bounced: bool


@dataclass(frozen=True, slots=True)
class WorkbookCampaign:
    """One tab: the campaign's header and its lists. ``index`` is the tab's position."""

    index: int
    name: str
    subject: str | None
    started_on: date
    last_batch_on: date | None
    recipients_count: int | None
    opens_count: int | None
    bounces_count: int | None
    opened: tuple[Person, ...]
    clicked: tuple[Person, ...]
    bounced: tuple[Person, ...]
    warnings: tuple[str, ...] = ()

    @property
    def recipients(self) -> tuple[Recipient, ...]:
        """Everyone any list names, once each by address, in the order first seen."""
        order: list[str] = []
        names: dict[str, Person] = {}
        flags: dict[str, set[ListKind]] = {}
        for kind, people in (
            (ListKind.OPENED, self.opened),
            (ListKind.CLICKED, self.clicked),
            (ListKind.BOUNCED, self.bounced),
        ):
            for person in people:
                if person.email not in names:
                    order.append(person.email)
                    names[person.email] = person
                flags.setdefault(person.email, set()).add(kind)
        return tuple(
            Recipient(
                first_name=names[email].first_name,
                last_name=names[email].last_name,
                email=email,
                opened=ListKind.OPENED in flags[email],
                clicked=ListKind.CLICKED in flags[email],
                bounced=ListKind.BOUNCED in flags[email],
            )
            for email in order
        )

    @property
    def unlisted(self) -> int | None:
        """How many recipients the ``Recipients`` count has that no list names."""
        if self.recipients_count is None:
            return None
        return max(self.recipients_count - len(self.recipients), 0)


@dataclass(frozen=True, slots=True)
class SkippedTab:
    index: int
    reason: str


@dataclass(frozen=True, slots=True)
class Workbook:
    sha256: str
    campaigns: tuple[WorkbookCampaign, ...]
    skipped: tuple[SkippedTab, ...] = field(default=())


# --- reading -----------------------------------------------------------------------


def read_workbook(data: bytes) -> Workbook:
    """Every campaign tab of the workbook in ``data``. :class:`WorkbookError` when it is
    not an xlsx file. A tab with no campaign header is skipped, with a reason."""
    digest = hashlib.sha256(data).hexdigest()
    _guard_zip(data)
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of types for a file it cannot read
        raise WorkbookError("not an xlsx workbook this reader understands") from exc
    campaigns: list[WorkbookCampaign] = []
    skipped: list[SkippedTab] = []
    try:
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        book.close()
        raise WorkbookError("not an xlsx workbook this reader understands") from exc
    try:
        for index, (sheet, formula_sheet) in enumerate(
            zip(book.worksheets, formulas.worksheets, strict=True)
        ):
            try:
                rows, uncached = _read_sheet(formula_sheet, sheet)
            except WorkbookError as exc:
                raise WorkbookError(f"tab {index + 1}: {exc}") from exc
            try:
                campaigns.append(
                    read_tab(index, rows, title=sheet.title, uncached_formulas=uncached)
                )
            except WorkbookError as exc:
                skipped.append(SkippedTab(index, str(exc)))
                log.info("history workbook: tab %d skipped (%s)", index + 1, exc)
    finally:
        book.close()
        formulas.close()
    return Workbook(sha256=digest, campaigns=tuple(campaigns), skipped=tuple(skipped))


def _guard_zip(data: bytes) -> None:
    """Refuse a zip bomb on its declared member sizes, before openpyxl opens it. The
    limits are the LinkedIn archive importer's. A file that is not a zip is left for
    ``load_workbook`` to refuse."""
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile:
        return
    try:
        guard_zip_members(infos, "workbook")
    except ZipRefusal as exc:
        raise WorkbookError(f"refused: {exc}") from exc


def _check_dimension(sheet: Any) -> None:
    rows, cols = sheet.max_row or 0, sheet.max_column or 0
    if rows > MAX_SHEET_ROWS:
        raise WorkbookError(f"refused: {rows} rows declared, over the {MAX_SHEET_ROWS} row limit")
    if cols > MAX_SHEET_COLUMNS:
        raise WorkbookError(
            f"refused: {cols} columns declared, over the {MAX_SHEET_COLUMNS} column limit"
        )
    if rows * cols > MAX_SHEET_CELLS:
        raise WorkbookError(
            f"refused: {rows} x {cols} cells declared, over the {MAX_SHEET_CELLS} cell limit"
        )


def _read_sheet(formulas: Any, values: Any) -> tuple[list[list[object]], int]:
    """One tab's rows (saved values) and how many formula cells have no saved value. Read with
    ``data_only``, such a cell is empty, which would silently drop a row; the tab reports it
    instead. One bounded pass: a sheet over the row, column or cell limit is refused as soon
    as a row crosses it, however little the file declared."""
    _check_dimension(values)
    rows: list[list[object]] = []
    cells = uncached = 0
    for formula_row, value_row in zip(
        formulas.iter_rows(values_only=True), values.iter_rows(values_only=True), strict=False
    ):
        if len(rows) >= MAX_SHEET_ROWS:
            raise WorkbookError(f"refused: the sheet has more than {MAX_SHEET_ROWS} rows")
        if len(value_row) > MAX_SHEET_COLUMNS:
            raise WorkbookError(
                f"refused: a row has {len(value_row)} columns, over the {MAX_SHEET_COLUMNS} limit"
            )
        cells += len(value_row)
        if cells > MAX_SHEET_CELLS:
            raise WorkbookError(f"refused: the sheet has more than {MAX_SHEET_CELLS} cells")
        rows.append(list(value_row))
        for formula, value in zip(formula_row, value_row, strict=False):
            if value is None and _is_formula(formula):
                uncached += 1
    return rows, uncached


def _is_formula(value: object) -> bool:
    if isinstance(value, str):
        return value.startswith("=")
    return value is not None and "formula" in type(value).__name__.lower()


def read_tab(
    index: int,
    rows: Sequence[Sequence[object]],
    *,
    title: str = "",
    uncached_formulas: int = 0,
) -> WorkbookCampaign:
    """One tab's campaign from its cell values, row by row. :class:`WorkbookError` when it
    has no campaign header or no start date."""
    grid = _Grid(rows)
    header = grid.find("campaign name")
    if header is None:
        raise WorkbookError("no 'Campaign name' header")
    top, _ = header
    labels = {col: label for col, label in grid.labels(top)}
    warnings: list[str] = []
    if uncached_formulas:
        warnings.append(
            f"{uncached_formulas} formula cell(s) have no saved value and read as empty;"
            " open the sheet, let it calculate, and export it again"
        )

    def value(match: str) -> object:
        for col, label in labels.items():
            if match in label:
                return grid.at(top + 1, col)
        return None

    name = _text(value("campaign name")) or title.strip()
    if not name:
        raise WorkbookError("no campaign name")
    started = _date(value("start"))
    if started is None:
        raise WorkbookError("no start date")
    subject = _text(value("subject"))
    if subject is not None and ("{{" in subject or "}}" in subject):
        warnings.append(
            "the subject is personalized ({{...}}), so the Gmail scan's exact subject"
            " match cannot find replies to it"
        )
    last_batch = _date(value("last batch"))
    if value("last batch") is not None and last_batch is None:
        warnings.append("the last batch date could not be read")

    lists: dict[ListKind, tuple[Person, ...]] = {}
    clicks_label = grid.find("clicks")
    for row, col in grid.find_all("first name"):
        kind = _list_kind(grid, row, col, clicks_label)
        if kind is None:
            warnings.append("a list with an unexpected header was skipped")
            continue
        if kind in lists:
            warnings.append(f"a second {kind} list was skipped")
            continue
        people, dropped = _walk(grid, row, col, kind)
        lists[kind] = people
        if dropped:
            warnings.append(f"{dropped} row(s) of the {kind} list were dropped: no usable address")

    opens = _count(value("opens"))
    opened = lists.get(ListKind.OPENED, ())
    if ListKind.OPENED not in lists:
        warnings.append("no opens list")
    elif opens is not None and opens != len(opened):
        warnings.append(f"the opens list has {len(opened)} people but the opens cell says {opens}")
    bounces = _count(value("bounce"))
    bounced = lists.get(ListKind.BOUNCED, ())
    if bounces is not None and bounces != len(bounced):
        warnings.append(
            f"the bounce list has {len(bounced)} people but the bounced cell says {bounces}"
        )
    return WorkbookCampaign(
        index=index,
        name=name,
        subject=subject,
        started_on=started,
        last_batch_on=last_batch,
        recipients_count=_count(value("recipients")),
        opens_count=opens,
        bounces_count=bounces,
        opened=opened,
        clicked=lists.get(ListKind.CLICKED, ()),
        bounced=bounced,
        warnings=tuple(warnings),
    )


class _Grid:
    """Cell values by zero-based row and column; anything outside is ``None``."""

    def __init__(self, rows: Sequence[Sequence[object]]) -> None:
        self.rows = rows

    def at(self, row: int, col: int) -> object:
        if row < 0 or col < 0 or row >= len(self.rows) or col >= len(self.rows[row]):
            return None
        return self.rows[row][col]

    def label(self, row: int, col: int) -> str | None:
        return _label(self.at(row, col))

    def labels(self, row: int) -> Iterable[tuple[int, str]]:
        for col in range(len(self.rows[row])):
            label = self.label(row, col)
            if label:
                yield col, label

    def find_all(self, wanted: str) -> list[tuple[int, int]]:
        return [
            (row, col)
            for row in range(len(self.rows))
            for col in range(len(self.rows[row]))
            if self.label(row, col) == wanted
        ]

    def find(self, wanted: str) -> tuple[int, int] | None:
        found = self.find_all(wanted)
        return found[0] if found else None

    def blank(self, row: int, cols: Iterable[int]) -> bool:
        return all(_text(self.at(row, col)) is None for col in cols)


def _list_kind(
    grid: _Grid, row: int, col: int, clicks_label: tuple[int, int] | None
) -> ListKind | None:
    if grid.label(row, col + 1) != "last name":
        return None
    third = grid.label(row, col + 2) or ""
    if "email" not in third and "bounce" not in third:
        return None
    if clicks_label == (row - 1, col):
        return ListKind.CLICKED
    if "bounce" in third:
        return ListKind.BOUNCED
    return ListKind.OPENED


def _walk(grid: _Grid, header_row: int, col: int, kind: ListKind) -> tuple[tuple[Person, ...], int]:
    """The people under a list header, and how many other rows were dropped. In the
    clicks list, a row holding only a link in the address column is the link the person
    above clicked, and is not counted; any other row without an address is."""
    people: list[Person] = []
    seen: set[str] = set()
    dropped = 0
    blanks = 0
    row = header_row + 1
    cols = (col, col + 1, col + 2)
    while row < len(grid.rows) and blanks < BLANK_ROWS_END_A_LIST:
        if grid.blank(row, cols):
            blanks += 1
            row += 1
            continue
        blanks = 0
        first, last = _text(grid.at(row, col)), _text(grid.at(row, col + 1))
        email = _address(grid.at(row, col + 2))
        if email is not None:
            if email not in seen:
                seen.add(email)
                people.append(Person(first, last, email))
        elif not (
            kind is ListKind.CLICKED and first is None and last is None and _is_link(grid, row, col)
        ):
            dropped += 1
        row += 1
    return tuple(people), dropped


def _is_link(grid: _Grid, row: int, col: int) -> bool:
    """Whether the address column holds a clicked link: link-shaped, and no ``@`` (a
    malformed address such as ``ada@example.test;`` is reported, not taken for a link)."""
    text = _text(grid.at(row, col + 2)) or ""
    return "@" not in text and bool(_LINK.search(text))


def _label(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    text = _SPACES.sub(" ", value.replace("\xa0", " ")).strip().lower().rstrip(":").strip()
    return text or None


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = _SPACES.sub(" ", str(value).replace("\xa0", " ")).strip()
    return text or None


def _address(value: object) -> str | None:
    text = _text(value)
    if text is None or not _ADDRESS.fullmatch(text):
        return None
    return text.lower()


def _count(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return int(value)
    if isinstance(value, str):
        match = _LEADING_INT.match(value)
        return int(match.group(1)) if match else None
    return None


def _date(value: object) -> date | None:
    """A cell as a date: a ``datetime`` or ``date``, an Excel serial number, or ISO text."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        try:
            converted = from_excel(value)
        except (ValueError, OverflowError):
            return None
        return converted.date() if isinstance(converted, datetime) else None
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None
